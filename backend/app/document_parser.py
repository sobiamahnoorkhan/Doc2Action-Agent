import io, os
import fitz
from docx import Document
from openpyxl import load_workbook

def extract_text(filename: str, data: bytes) -> str:
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        doc = fitz.open(stream=data, filetype="pdf")
        return "\n".join(page.get_text() for page in doc)
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    if ext == ".xlsx":
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                out.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(out)
    if ext in {".txt", ".md"}:
        return data.decode("utf-8", errors="replace")
    raise ValueError("Supported formats: PDF, DOCX, XLSX, TXT, MD")
